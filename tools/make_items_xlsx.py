#!/usr/bin/env python3
"""Build the programmer's open-items workbooks from docs/open_items.json.

    python3 tools/make_items_xlsx.py [--scope gestamp|mejoras|both] [--out FILE.xlsx]

Two workbooks since 2026-10-12 (tools/scope.py): the Gestamp points G01-G22,
which the customer audits for the buy-off, and the Ethos items as process
improvements. By default it writes both to dist/
(R10_Gestamp_Puntos_<YYYY-MM-DD_HHMM>.xlsx and R10_Mejoras_Puntos_<...>.xlsx,
date and time of central Mexico, tools/naming.py) and removes the older
copies in dist/. --out needs a single --scope.

docs/open_items.json is the master list (kept in the repository and updated
after every audit). The workbook is the working copy the robot programmer
fills in during the day and sends back with the robot backup:

  Puntos abiertos   one row per item. Yellow columns are the programmer's;
                    grey columns are written by the audit.
  Resumen           counts by status and severity (formulas).
  Instrucciones     how to fill it in, with one example row.
  Historial         one row per audit.

tools/audit_backup.py --items reads the 'Puntos abiertos' sheet by header
text, so the header names below are part of the interface.
"""
import argparse
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import scope as scopes  # noqa: E402

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# 'Corregido en oficina - probar en celda' was dropped on 2026-10-11: the programmer's program is the base.
STATES = ['Abierto', 'Requiere decisión', 'En proceso',
          'Corregido - por auditar', 'Cerrado - auditado', 'Reabierto', 'Aceptado sin cambio',
          'Cerrado - depuración', 'No aplica', 'Sin acción - documentado', 'Se cierra con otro punto']
PENDING = ['Abierto', 'Requiere decisión', 'En proceso',
           'Corregido - por auditar', 'Reabierto']
PROGRAMMER_STATES = ['En proceso', 'Corregido - por auditar']
# Closed: nothing is left to do. These rows are green, critical or not.
CLOSED = ['Cerrado - auditado', 'Cerrado - depuración', 'Aceptado sin cambio', 'No aplica',
          'Sin acción - documentado', 'Se cierra con otro punto']
SEVERITIES = ['Crítico', 'Alto', 'Medio', 'Bajo', '-']

# (header, width, owner) - owner: 'item' fixed text, 'prog' programmer input, 'audit' auditor.
COLUMNS = [
    ('ID', 6, 'item'), ('Origen', 11, 'item'), ('Sev.', 8, 'item'), ('Punto', 46, 'item'),
    ('Tipo', 18, 'item'),
    ('Responsable', 18, 'item'), ('Módulos', 28, 'item'), ('Criterio de cierre', 70, 'item'),
    ('Depende de', 12, 'item'),
    ('Estado', 22, 'prog'), ('Qué cambió el programador', 50, 'prog'),
    ('Módulos modificados', 28, 'prog'), ('Evidencia / prueba en celda', 40, 'prog'),
    ('Fecha programador', 13, 'prog'),
    ('Resultado auditoría', 20, 'audit'), ('Fecha auditoría', 13, 'audit'),
    ('Nota auditoría', 50, 'audit'),
]


def columns(scope):
    """The Gestamp list names the related Ethos items with their state; the improvement list
    says which open Gestamp point each item counts for."""
    if scope == 'gestamp':
        return [('Relacionados (Ethos)', 24, 'item') if h == 'Depende de' else (h, w, o) for h, w, o in COLUMNS]
    if scope == 'mejoras':
        k = [h for h, _, _ in COLUMNS].index('Depende de') + 1
        return COLUMNS[:k] + [('Cuenta para Gestamp', 12, 'item')] + COLUMNS[k:]
    return COLUMNS

FONT = 'Arial'
THIN = Side(style='thin', color='CBD2D9')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill('solid', fgColor='1D4E89')
PROG = PatternFill('solid', fgColor='FFF4C2')
AUDIT = PatternFill('solid', fgColor='E4E7EB')
SEV_COLOR = {'Crítico': 'B42318', 'Alto': 'C4561D', 'Medio': 'A16207', 'Bajo': '52606D', '-': '9AA5B1'}
GESTAMP = PatternFill('solid', fgColor='DCEBFA')
DONE = PatternFill('solid', fgColor='C6EFCE')


def is_gestamp(item):
    return item['id'].startswith('G')


def origin_text(item):
    if is_gestamp(item):
        return f"Gestamp #{int(item['id'][1:])}"
    return item.get('origen_ref') or 'Ethos'


def title_text(item):
    if item.get('texto_gestamp'):
        return f"{item['titulo']}\nGestamp: \"{item['texto_gestamp']}\""
    return item['titulo']


def criteria_text(item, by_id=None):
    """Bullets; with by_id (Gestamp list) a 'se cumplen los criterios de F43' carries F43's criteria."""
    crit = scopes.expanded_criteria(item, by_id) if by_id else [(False, c) for c in item.get('criterio_cierre', [])]
    return '\n'.join(('    – ' if sub else '• ') + c for sub, c in crit)


def items_sheet(ws, items, scope='both', all_items=None):
    ws.title = 'Puntos abiertos'
    cols = columns(scope)
    by_id = {i['id']: i for i in (all_items or items)}
    links = scopes.gestamp_links(all_items or items)
    for col, (head, width, owner) in enumerate(cols, 1):
        c = ws.cell(row=1, column=col, value=head)
        c.font = Font(name=FONT, bold=True, color='FFFFFF', size=9)
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical='center')
        c.border = BORDER
        ws.column_dimensions[get_column_letter(col)].width = width
    status_col = [h for h, _, _ in cols].index('Estado') + 1
    ws.cell(row=1, column=status_col).comment = Comment(
        'El programador solo pone En proceso o Corregido - por auditar. '
        'Cerrado lo pone la auditoría.', 'Ethos')
    for r, it in enumerate(items, 2):
        audit = (it.get('historial') or [{}])[-1]
        gest = scope == 'gestamp'
        cell = {'ID': it['id'], 'Origen': origin_text(it), 'Sev.': it['severidad'], 'Punto': title_text(it),
                'Tipo': it['tipo'], 'Responsable': it['responsable'], 'Módulos': ', '.join(it.get('modulos', [])),
                'Criterio de cierre': criteria_text(it, by_id if gest else None),
                'Depende de': ', '.join(it.get('depende_de', [])),
                'Relacionados (Ethos)': scopes.with_state(scopes.related(it, by_id), by_id),
                'Cuenta para Gestamp': ', '.join(links.get(it['id'], [])),
                'Estado': it.get('estado', 'Abierto'), 'Qué cambió el programador': it.get('nota_programador', ''),
                'Módulos modificados': it.get('modulos_modificados', ''),
                'Evidencia / prueba en celda': it.get('evidencia', ''),
                'Fecha programador': it.get('fecha_programador', ''),
                'Resultado auditoría': audit.get('resultado', ''), 'Fecha auditoría': audit.get('fecha', ''),
                'Nota auditoría': audit.get('nota', '')}
        values = [cell[h] for h, _, _ in cols]
        for col, ((head, _, owner), v) in enumerate(zip(cols, values), 1):
            c = ws.cell(row=r, column=col, value=v)
            c.font = Font(name=FONT, size=9, bold=(head == 'ID'),
                          color=SEV_COLOR.get(v, '1F2933') if head == 'Sev.' else '1F2933')
            c.alignment = Alignment(wrap_text=True, vertical='top')
            c.border = BORDER
            if it.get('estado') in CLOSED:
                c.fill = DONE
            elif owner == 'prog':
                c.fill = PROG
            elif owner == 'audit':
                c.fill = AUDIT
            elif (is_gestamp(it) or links.get(it['id'])) and head in ('ID', 'Origen', 'Punto',
                                                                      'Cuenta para Gestamp'):
                c.fill = GESTAMP
    last = len(items) + 1
    dv = DataValidation(type='list', formula1='"' + ','.join(STATES) + '"', allow_blank=False,
                        showErrorMessage=True, errorTitle='Estado',
                        error='Elige un estado de la lista (hoja Instrucciones).')
    ws.add_data_validation(dv)
    status = get_column_letter(status_col)
    dv.add(f'{status}2:{status}{last}')
    # The row also turns green when the auditor sets a closed state in the workbook.
    cond = ' , '.join(f'${status}2="{st}"' for st in CLOSED).replace(' , ', ',')
    ws.conditional_formatting.add(f'A2:{get_column_letter(len(cols))}{last}',
                                  FormulaRule(formula=[f'OR({cond})'], fill=DONE, stopIfTrue=True))
    ws.freeze_panes = 'D2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{last}'
    ws.sheet_view.zoomScale = 90


def summary_sheet(ws, n, status_col, generated, scope='both'):
    ws.title = 'Resumen'
    cols = columns(scope)
    ws['A1'] = {'gestamp': 'Resumen de los puntos de Gestamp (buy off)',
                'mejoras': 'Resumen de las mejoras de proceso (puntos de Ethos)'}.get(scope, 'Resumen de puntos abiertos')
    ws['A1'].font = Font(name=FONT, bold=True, size=12, color='1D4E89')
    ws['A2'] = f'Generado {generated}'
    ws['A2'].font = Font(name=FONT, size=8, italic=True, color='616E7C')
    ws['A3'], ws['B3'] = 'Estado', 'Puntos'
    for col in ('A3', 'B3'):
        ws[col].font = Font(name=FONT, bold=True, size=9)
        ws[col].fill = AUDIT
    col = get_column_letter(status_col)
    rng = f"'Puntos abiertos'!${col}$2:${col}${n + 1}"
    for r, s in enumerate(STATES, 4):
        ws.cell(row=r, column=1, value=s).font = Font(name=FONT, size=9)
        c = ws.cell(row=r, column=2, value=f'=COUNTIF({rng},A{r})')
        c.font = Font(name=FONT, size=9)
    tot = 4 + len(STATES)
    ws.cell(row=tot, column=1, value='Total').font = Font(name=FONT, bold=True, size=9)
    ws.cell(row=tot, column=2, value=f'=SUM(B4:B{tot - 1})').font = Font(name=FONT, bold=True, size=9)

    top = tot + 2
    ws.cell(row=top, column=1, value='Severidad').font = Font(name=FONT, bold=True, size=9)
    ws.cell(row=top, column=2, value='Total').font = Font(name=FONT, bold=True, size=9)
    ws.cell(row=top, column=3, value='Pendientes').font = Font(name=FONT, bold=True, size=9)
    for col in range(1, 4):
        ws.cell(row=top, column=col).fill = AUDIT
    sev_col = get_column_letter([h for h, _, _ in cols].index('Sev.') + 1)
    sev = f"'Puntos abiertos'!${sev_col}$2:${sev_col}${n + 1}"
    for r, s in enumerate(SEVERITIES[:-1], top + 1):
        ws.cell(row=r, column=1, value=s).font = Font(name=FONT, size=9)
        ws.cell(row=r, column=2, value=f'=COUNTIF({sev},A{r})').font = Font(name=FONT, size=9)
        parts = '+'.join(f'COUNTIFS({sev},A{r},{rng},"{p}")' for p in PENDING)
        ws.cell(row=r, column=3, value='=' + parts).font = Font(name=FONT, size=9)
    org_col = get_column_letter([h for h, _, _ in cols].index('Origen') + 1)
    org = f"'Puntos abiertos'!${org_col}$2:${org_col}${n + 1}"
    gtop = top + len(SEVERITIES) + 1
    ws.cell(row=gtop, column=1, value='Origen').font = Font(name=FONT, bold=True, size=9)
    ws.cell(row=gtop, column=2, value='Total').font = Font(name=FONT, bold=True, size=9)
    ws.cell(row=gtop, column=3, value='Pendientes').font = Font(name=FONT, bold=True, size=9)
    for col in range(1, 4):
        ws.cell(row=gtop, column=col).fill = AUDIT
    origins = [('Gestamp', 'Gestamp*'), ('Ethos', 'Ethos'), ('Calvin (Ethos)', 'Calvin*')]
    origins = {'gestamp': origins[:1], 'mejoras': origins[1:]}.get(scope, origins)
    for r, (label, pat) in enumerate(origins, gtop + 1):
        ws.cell(row=r, column=1, value=label).font = Font(name=FONT, size=9)
        ws.cell(row=r, column=2, value=f'=COUNTIF({org},"{pat}")').font = Font(name=FONT, size=9)
        parts = '+'.join(f'COUNTIFS({org},"{pat}",{rng},"{p}")' for p in PENDING)
        ws.cell(row=r, column=3, value='=' + parts).font = Font(name=FONT, size=9)
    note = gtop + 5
    ws.cell(row=note, column=1, value='Pendientes = ' + ', '.join(PENDING) + '.').font = \
        Font(name=FONT, size=8, italic=True, color='616E7C')
    ws.column_dimensions['A'].width = 34
    ws.column_dimensions['B'].width = 10
    ws.column_dimensions['C'].width = 12


SCOPE_NOTE = {
    'gestamp': [
        ('Esta lista: los puntos de Gestamp (buy off)', True),
        ('Son los 22 puntos de la revisión de Gestamp del 2026-10-02 (G01–G22): los que Gestamp audita para el buy '
         'off. Van primero. Cuando un criterio dice "se cumplen los criterios de F43", los criterios de F43 vienen '
         'debajo, con guion: F43 está en la lista de mejoras y se cierra junto con este punto. La columna '
         '"Relacionados (Ethos)" da los puntos de Ethos ligados, con su estado.', False),
        ('', False)],
    'mejoras': [
        ('Esta lista: mejoras de proceso (puntos de Ethos)', True),
        ('Son los puntos de la revisión de Ethos (Fnn/Cnn). No condicionan el buy off: se trabajan después de los '
         'de Gestamp, cuando haya tiempo. Si la columna "Cuenta para Gestamp" trae un punto G, cerrarlo también '
         'ayuda a cerrar ese punto de Gestamp (está en la otra lista).', False),
        ('', False)],
}


def instructions_sheet(ws, scope='both'):
    ws.title = 'Instrucciones'
    name = scopes.SCOPES[scope]['xlsx'] if scope in scopes.SCOPES else 'R10_Puntos_Abiertos'
    rows = SCOPE_NOTE.get(scope, []) + [
        ('Cómo usar esta lista', True),
        ('1. Trabaja los puntos en el programa del robot. Respeta docs/CONVENTIONS.md del repositorio: '
         'comentarios en inglés con la señal nombrada ($OUT[n] nombre), sin código comentado, sin reabrir '
         'inline forms que tengan líneas escritas a mano.', False),
        ('2. En la hoja "Puntos abiertos" llena SOLO las columnas amarillas de los puntos que trabajaste: '
         'Estado, Qué cambió el programador, Módulos modificados, Evidencia / prueba en celda, Fecha '
         'programador.', False),
        ('3. Estado: pon "En proceso" si empezaste y no está listo, o "Corregido - por auditar" si el '
         'cambio está hecho y probado. '
         'No pongas Cerrado: lo pone la auditoría.', False),
        ('4. La evidencia es lo que pide el criterio de cierre: prueba en celda con fecha y resultado, '
         'valor leído en el smartPAD, documento de Gestamp. Sin evidencia el punto se reabre.', False),
        ('5. Al final del día saca un respaldo completo del robot (Archive → All), nómbralo '
         '658424_AAAA-MM-DD_HHMM.zip con la hora en que lo sacaste, y envíalo junto con este archivo '
         f'guardado como {name}_AAAA-MM-DD_HHMM.xlsx (y el otro Excel si trabajaste puntos de esa lista). '
         'No borres filas ni cambies las columnas blancas o grises.', False),
        ('', False),
        ('Ejemplo de una fila llenada (formato esperado)', True),
        ('Estado: Corregido - por auditar', False),
        ('Qué cambió el programador: style1app1opt1..3: espera supervisada antes del LIN de aproximación '
         '(gun open LPT > 120000, $IN[466], $IN[482], NOT $IN[470]) con $TIMER[21] 3 s, mensaje y '
         'do004ProcessFault. PRELOAD paso 20 ahora espera $IN[482].', False),
        ('Módulos modificados: style1app1opt1.src ... style1app1opt3.src, PRELOAD.src', False),
        ('Evidencia / prueba en celda: 2026-10-06, T1 50 %: clamp cerrado a mano, el robot se detuvo '
         'antes de P22 con mensaje; con todo abierto el ciclo siguió normal. 3 ciclos en automático OK.', False),
        ('Fecha programador: 2026-10-06', False),
        ('', False),
        ('Puntos de Gestamp', True),
        ('Los puntos G01 a G22 vienen de la revisión de Gestamp del 2026-10-02 (BMW-03-10R1): columna Origen '
         '"Gestamp #n", fondo azul y el texto original de Gestamp en la columna Punto. Los Fnn/Cnn son de la revisión '
         'de Ethos y conservan la numeración del R20 (F37 no aplica en el R10; F44 a F52 son solo del R10). La lista '
         'de Calvin se verificó punto por punto: los puntos que él también señaló dicen "Calvin #n" en Origen. Esta '
         'lista es la única que se usa.', False),
        ('Desde el 2026-10-11 la base es tu programa: no se carga ninguna versión de oficina. Cada respaldo se '
         'audita contra el anterior. En los puntos que la oficina había corregido, la Nota auditoría dice qué le '
         'falta a tu programa; el PDF y docs/OPEN_ITEMS.es.md traen qué hacer y, como referencia, cómo lo resolvió '
         'la oficina.', False),
        ('Fila completa en verde: punto cerrado (Cerrado, Aceptado sin cambio, No aplica, Sin acción o Se cierra '
         'con otro punto), aunque sea crítico. No hay que hacer nada en esas filas.', False),
        ('Las anotaciones que Edgar acuerda con el programador van en la columna "Qué cambió el programador".', False),
        ('Estándar: no se pide más de lo que hace el respaldo de Gestamp NUT07_10R1 (ver la Nota auditoría de cada '
         'punto).', False),
        ('', False),
        ('Estados', True),
    ]
    meaning = {
        'Abierto': 'Pendiente de trabajar.',
        'Requiere decisión': 'Falta información o decisión (Gestamp, seguridad, Ethos) antes de programar.',
        'En proceso': 'Lo pone el programador: trabajo empezado.',
        'Corregido - por auditar': 'Lo pone el programador: cambio hecho y probado, pide auditoría.',
        'Cerrado - auditado': 'Lo pone la auditoría: cumple el criterio de cierre en el respaldo.',
        'Reabierto': 'Lo pone la auditoría: no se encontró el cambio o no cumple (ver Nota auditoría).',
        'Aceptado sin cambio': 'Ethos / Gestamp decidieron no corregir; queda la razón.',
        'Cerrado - depuración': 'Resuelto por la depuración del 2026-10-04.',
        'No aplica': 'El punto ya cumple.',
        'Sin acción - documentado': 'Se deja como está; la nota dice cuándo se reabriría.',
        'Se cierra con otro punto': 'Duplica otro punto (columna Depende de); se cierra con él.',
    }
    r = 1
    for text, bold in rows:
        c = ws.cell(row=r, column=1, value=text)
        c.font = Font(name=FONT, size=10 if bold else 9, bold=bold, color='1D4E89' if bold else '1F2933')
        c.alignment = Alignment(wrap_text=True, vertical='top')
        r += 1
    for s in STATES:
        ws.cell(row=r, column=1, value=f'{s}: {meaning[s]}').font = Font(name=FONT, size=9)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value='Colores: amarillo = lo llena el programador; gris = lo llena la '
                                   'auditoría; blanco = no se edita.').font = Font(name=FONT, size=9,
                                                                                    italic=True)
    ws.column_dimensions['A'].width = 120


def history_sheet(ws, audits):
    ws.title = 'Historial'
    heads = ['Fecha auditoría', 'Respaldo', 'Cerrados', 'Reabiertos', 'Cambios de código', 'Notas']
    for col, h in enumerate(heads, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = Font(name=FONT, bold=True, size=9, color='FFFFFF')
        c.fill = HEAD
    for r, a in enumerate(audits, 2):
        for col, key in enumerate(['fecha', 'respaldo', 'cerrados', 'reabiertos', 'cambios', 'notas'], 1):
            ws.cell(row=r, column=col, value=a.get(key, '')).font = Font(name=FONT, size=9)
    for col, w in enumerate([14, 30, 10, 10, 16, 80], 1):
        ws.column_dimensions[get_column_letter(col)].width = w


def build(data, scope, out, when):
    import naming
    all_items = data['items']
    items = scopes.select(all_items, scope) if scope in scopes.SCOPES else all_items
    wb = Workbook()
    items_sheet(wb.active, items, scope, all_items)
    summary_sheet(wb.create_sheet(), len(items), [h for h, _, _ in columns(scope)].index('Estado') + 1,
                  naming.human(when), scope)
    instructions_sheet(wb.create_sheet(), scope)
    history_sheet(wb.create_sheet(), data.get('auditorias', []))
    # The summary sheet is formulas only: have Excel compute them when the file opens.
    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    wb.save(out)
    print(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--items', default=os.path.join(REPO, 'docs', 'open_items.json'))
    ap.add_argument('--scope', choices=['gestamp', 'mejoras', 'both', 'all'], default='both',
                    help="gestamp, mejoras, both (default: the two workbooks) or all (one workbook, every item)")
    ap.add_argument('--out', help='output file (default: dated file in dist/); needs a single --scope')
    args = ap.parse_args()
    if args.out and args.scope == 'both':
        ap.error('--out needs --scope gestamp, mejoras or all')
    import naming
    when = naming.now()
    data = json.load(open(args.items))
    if args.out:
        build(data, args.scope, args.out, when)
        return
    dist = os.path.join(REPO, 'dist')
    for scope in (['gestamp', 'mejoras'] if args.scope == 'both' else [args.scope]):
        prefix = scopes.SCOPES[scope]['xlsx'] if scope in scopes.SCOPES else scopes.OLD['xlsx']
        out = naming.dated_path(dist, prefix, '.xlsx', when)
        build(data, scope, out, when)
        for old in naming.remove_older(out, prefix, '.xlsx'):
            print('removed', old)
    if args.scope == 'both':
        # the single workbook of before 2026-10-12
        for old in naming.remove_older(os.path.join(dist, 'x.xlsx'), scopes.OLD['xlsx'], '.xlsx'):
            print('removed', old)


if __name__ == '__main__':
    main()
